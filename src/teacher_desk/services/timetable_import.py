"""Read common timetable files and suggest weekly entries for teacher review."""

from __future__ import annotations

import csv
from bisect import bisect_right
from collections import Counter, defaultdict
import io
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from teacher_desk.services.timetable_vision import (
    DEFAULT_MODEL,
    DEFAULT_OPENROUTER_MODEL,
    recognize_image,
    recognize_image_openrouter,
    recognize_image_nvidia,
)
from teacher_desk.services.word_convert import which_soffice


WEEKDAYS = ("星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日")
DAY_ALIASES = {
    "一": 0, "二": 1, "三": 2, "四": 3, "五": 4, "六": 5, "日": 6, "天": 6,
    "mon": 0, "monday": 0, "tue": 1, "tuesday": 1, "wed": 2, "wednesday": 2,
    "thu": 3, "thursday": 3, "fri": 4, "friday": 4, "sat": 5, "saturday": 5,
    "sun": 6, "sunday": 6,
    "mo": 0, "tu": 1, "we": 2, "th": 3, "fr": 4, "sa": 5, "su": 6,
}
DAY_PATTERN = re.compile(
    r"星期[一二三四五六日天]|[週周][一二三四五六日天]|"
    r"\b(?:monday|tuesday|wednesday|thursday|friday|saturday|sunday|mon|tue|wed|thu|fri|sat|sun|mo|tu|we|th|fr|sa|su)\b",
    re.IGNORECASE,
)
TIME_PATTERN = re.compile(r"(?<!\d)([0-2]?\d)\s*[:：.]\s*([0-5]\d)(?!\d)")
SPAN_PATTERN = re.compile(
    r"(?<!\d)([0-2]?\d\s*[:：.]\s*[0-5]\d)\s*(?:-|–|—|~|～|至|到)\s*"
    r"([0-2]?\d\s*[:：.]\s*[0-5]\d)(?!\d)"
)
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


@dataclass(frozen=True)
class SuggestedEntry:
    weekday: int
    start_time: str
    end_time: str
    title: str
    location: str = ""
    enabled: bool = True


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (time, datetime)):
        return value.strftime("%H:%M")
    return str(value).strip()


def _rows_text(rows: list[list[object]]) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    writer.writerows([_cell_text(cell) for cell in row] for row in rows)
    return output.getvalue().rstrip("\n")


def _read_excel(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return "\n\n".join(
            _rows_text([list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    blocks = [
        _rows_text([[cell.text for cell in row.cells] for row in table.rows])
        for table in document.tables
    ]
    blocks.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    return "\n\n".join(blocks)


def _read_doc(path: Path) -> str:
    soffice = which_soffice()
    if soffice:
        with tempfile.TemporaryDirectory() as directory:
            result = subprocess.run(
                [str(soffice), "--headless", "--convert-to", "docx", "--outdir", directory, str(path)],
                capture_output=True, text=True, timeout=120, check=False,
            )
            converted = Path(directory) / f"{path.stem}.docx"
            if result.returncode == 0 and converted.exists():
                return _read_docx(converted)
    if sys.platform == "win32":
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore

            pythoncom.CoInitialize()
            try:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.AutomationSecurity = 3
                try:
                    document = word.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
                    with tempfile.TemporaryDirectory() as directory:
                        converted = Path(directory) / f"{path.stem}.docx"
                        try:
                            document.SaveAs2(str(converted), FileFormat=16)
                        finally:
                            document.Close(False)
                        return _read_docx(converted)
                finally:
                    word.Quit()
            finally:
                pythoncom.CoUninitialize()
        except (ImportError, OSError) as exc:
            raise RuntimeError("舊版 .doc 需要 Microsoft Word 或 LibreOffice 才能讀取") from exc
    raise RuntimeError("舊版 .doc 需要 Microsoft Word 或 LibreOffice 才能讀取")


def _read_xls(path: Path) -> str:
    soffice = which_soffice()
    if soffice:
        with tempfile.TemporaryDirectory() as directory:
            result = subprocess.run(
                [str(soffice), "--headless", "--convert-to", "xlsx", "--outdir", directory, str(path)],
                capture_output=True, text=True, timeout=120, check=False,
            )
            converted = Path(directory) / f"{path.stem}.xlsx"
            if result.returncode == 0 and converted.exists():
                return _read_excel(converted)
    if sys.platform == "win32":
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore

            pythoncom.CoInitialize()
            try:
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                excel.AutomationSecurity = 3
                try:
                    workbook = excel.Workbooks.Open(str(path.resolve()), ReadOnly=True)
                    with tempfile.TemporaryDirectory() as directory:
                        converted = Path(directory) / f"{path.stem}.xlsx"
                        try:
                            workbook.SaveAs(str(converted), FileFormat=51)
                        finally:
                            workbook.Close(False)
                        return _read_excel(converted)
                finally:
                    excel.Quit()
            finally:
                pythoncom.CoUninitialize()
        except (ImportError, OSError) as exc:
            raise RuntimeError("舊版 .xls 需要 Microsoft Excel 或 LibreOffice 才能讀取") from exc
    raise RuntimeError("舊版 .xls 需要 Microsoft Excel 或 LibreOffice 才能讀取")


def find_tesseract() -> Path | None:
    bundle_root = getattr(sys, "_MEIPASS", None)
    if bundle_root:
        bundled = Path(bundle_root) / "tesseract" / "tesseract.exe"
        if bundled.is_file():
            return bundled
    found = shutil.which("tesseract")
    if found:
        return Path(found)
    for path in (
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ):
        if path.is_file():
            return path
    return None


def _run_tesseract(tesseract: Path, path: Path, languages: str, page_mode: int) -> str:
    result = subprocess.run(
        [str(tesseract), str(path), "stdout", "-l", languages, "--psm", str(page_mode)],
        capture_output=True, encoding="utf-8", errors="replace", timeout=120, check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or "OCR 讀取失敗")
    return result.stdout.strip()


def _grid_boundaries(path: Path) -> tuple[list[int], list[int]] | None:
    """Find long ruled lines so OCR can read timetable cells separately."""
    from PIL import Image

    with Image.open(path) as source:
        original_width, original_height = source.size
        image = source.convert("L")
        image.thumbnail((1800, 1800))
    width, height = image.size
    pixels = image.tobytes()
    column_counts = [0] * width
    row_counts = [0] * height
    for y in range(height):
        row = pixels[y * width:(y + 1) * width]
        for x, value in enumerate(row):
            if value < 150:
                row_counts[y] += 1
                column_counts[x] += 1

    def centers(counts: list[int], minimum: int) -> list[int]:
        found = [index for index, count in enumerate(counts) if count >= minimum]
        groups: list[list[int]] = []
        for index in found:
            if not groups or index > groups[-1][-1] + 1:
                groups.append([])
            groups[-1].append(index)
        return [round(sum(group) / len(group)) for group in groups]

    xs = centers(column_counts, round(height * 0.65))
    ys = centers(row_counts, round(width * 0.65))
    if not (3 <= len(xs) <= 15 and 3 <= len(ys) <= 30):
        return None
    xs = [round(x * original_width / width) for x in xs]
    ys = [round(y * original_height / height) for y in ys]
    if any(right - left < 35 for left, right in zip(xs, xs[1:])):
        return None
    if any(bottom - top < 30 for top, bottom in zip(ys, ys[1:])):
        return None
    return xs, ys


def _weekly_grid_boundaries(path: Path) -> tuple[list[int], list[int]] | None:
    """Find the time column and weekday columns of a ruled weekly timetable.

    The time column has every row line; lesson columns may omit a line when a
    lesson spans two periods, so their horizontal projections cannot be used.
    """
    from PIL import Image

    with Image.open(path) as source:
        image = source.convert("L")
    width, height = image.size
    pixels = image.tobytes()
    threshold = 210  # photographed light-grey grid lines are still visible

    def centers(counts: list[int], minimum: int) -> list[int]:
        runs: list[list[int]] = []
        for index, count in enumerate(counts):
            if count < minimum:
                continue
            if not runs or index > runs[-1][-1] + 1:
                runs.append([])
            runs[-1].append(index)
        return [round(sum(run) / len(run)) for run in runs]

    x_counts = [0] * width
    margin = round(width * 0.012)
    for y in range(height):
        row = pixels[y * width:(y + 1) * width]
        for x in range(margin, width - margin):
            if row[x] < threshold:
                x_counts[x] += 1
    xs = [x for x in centers(x_counts, round(height * 0.5)) if margin < x < width - margin]
    if not 4 <= len(xs) <= 9:
        return None
    if any(right - left < width * 0.04 for left, right in zip(xs, xs[1:])):
        return None

    # The first vertical rule is continuous from the table top to bottom.
    runs: list[list[int]] = []
    for y in range(height):
        if pixels[y * width + xs[0]] >= threshold:
            continue
        if not runs or y > runs[-1][-1] + 1:
            runs.append([])
        runs[-1].append(y)
    if not runs:
        return None
    table_run = max(runs, key=len)
    if len(table_run) < height * 0.4:
        return None
    left, right = xs[0] + 4, xs[1] - 4
    y_counts = [0] * height
    for y in range(table_run[0], table_run[-1] + 1):
        row = pixels[y * width:(y + 1) * width]
        y_counts[y] = sum(row[x] < threshold for x in range(left, right))
    ys = [y for y in centers(y_counts, round((right - left) * 0.75))
          if table_run[0] - 2 <= y <= table_run[-1] + 2]
    if not 5 <= len(ys) <= 30:
        return None
    if any(bottom - top < height * 0.015 for top, bottom in zip(ys, ys[1:])):
        return None
    return xs, ys


def _native_pdf_periods(text: str, expected: int) -> list[tuple[str, tuple[str, str]]] | None:
    """Use the PDF's precise printed times even when its Chinese font is broken."""
    lines = text.splitlines()
    periods = [
        (line.strip(), span)
        for index, line in enumerate(lines[:-1])
        if re.fullmatch(r"\d{1,2}|Recess\d+|Lunch", line.strip(), re.IGNORECASE)
        if (span := _span(lines[index + 1])) is not None
    ]
    return periods if len(periods) == expected else None


def _weekly_grid_text(
    path: Path, words: list[dict[str, object]] | None = None, native_text: str = "",
) -> str:
    """Build editable reminder rows from lesson cells, including merged periods."""
    bounds = _weekly_grid_boundaries(path)
    if bounds is None:
        return ""
    from PIL import Image

    xs, ys = bounds
    with Image.open(path) as source:
        image = source.convert("RGB")
    width, height = image.size
    grey = image.convert("L").tobytes()
    tesseract = find_tesseract()
    available = ""
    if tesseract:
        available = subprocess.run(
            [str(tesseract), "--list-langs"], capture_output=True, encoding="utf-8", errors="replace",
            timeout=15, check=False,
        ).stdout
    language_names = set(available.splitlines())
    header_language = "eng" if "eng" in language_names else None
    # A single Traditional Chinese model reads mixed Chinese/Latin aSc cells
    # more reliably than Tesseract's competing Chinese and English models.
    cell_languages = "chi_tra" if "chi_tra" in language_names else "eng" if "eng" in language_names else ""

    def words_in_box(left: int, top: int, right: int, bottom: int) -> str:
        if not words:
            return ""
        selected = [word for word in words if left + 2 < float(word["x"]) * width < right - 2
                    and top + 2 < float(word["y"]) * height < bottom - 2]
        lines: list[list[dict[str, object]]] = []
        for word in sorted(selected, key=lambda item: (float(item["y"]), float(item["x"]))):
            if not lines or float(word["y"]) - float(lines[-1][0]["y"]) > 0.015:
                lines.append([])
            lines[-1].append(word)
        return "\n".join(
            " ".join(str(word["text"]) for word in sorted(line, key=lambda item: float(item["x"])))
            for line in lines
        )

    with tempfile.TemporaryDirectory() as directory:
        crop_path = Path(directory) / "cell.png"

        def ocr_box(left: int, top: int, right: int, bottom: int, languages: str, psm: int) -> str:
            if not tesseract or not languages or right - left < 10 or bottom - top < 10:
                return ""
            crop = image.crop((left, top, right, bottom))
            factor = 3 if width < 1600 else 2 if width < 2400 else 1
            if factor > 1:
                crop = crop.resize((crop.width * factor, crop.height * factor), Image.Resampling.LANCZOS)
            crop.save(crop_path)
            return _run_tesseract(tesseract, crop_path, languages, psm)

        owner = ocr_box(round(width * 0.32), 4, round(width * 0.68), ys[0] - 4,
                        "chi_tra" if "chi_tra" in language_names else "", 7).strip()
        if not (2 <= len(owner) <= 8 and re.fullmatch(r"[\u3400-\u9fff]+", owner)):
            owner = ""

        weekdays: list[int | None] = []
        for column in range(1, len(xs) - 1):
            box = (xs[column] + 4, ys[0] + 3, xs[column + 1] - 4, ys[1] - 3)
            header = words_in_box(*box) or ocr_box(*box, header_language or "", 7)
            weekdays.append(_weekday(header))
        if sum(day is not None for day in weekdays) < 2:
            return ""

        period_bounds = ys[1:]
        periods = _native_pdf_periods(native_text, len(period_bounds) - 1) if native_text else None
        if periods is None:
            periods = []
            for top, bottom in zip(period_bounds, period_bounds[1:]):
                box = (xs[0] + 4, top + 3, xs[1] - 4, bottom - 3)
                period_text = ocr_box(*box, header_language or "", 6) or words_in_box(*box)
                first_line = period_text.splitlines()[0].strip() if period_text.splitlines() else ""
                label_match = re.match(r"^(\d{1,2}|Recess\s*\d+|Lunch)(?=\s|$)", first_line, re.IGNORECASE)
                periods.append((label_match.group() if label_match else "", _span(period_text)))

        rows: list[list[object]] = [["星期", "開始", "結束", "課堂", "地點", "啟用"]]
        boundary_index = {y: index for index, y in enumerate(period_bounds)}
        for column, weekday in enumerate(weekdays, start=1):
            if weekday is None:
                continue
            left, right = xs[column], xs[column + 1]
            column_bounds = [period_bounds[0]]
            for y in period_bounds[1:]:
                inner_left, inner_right = left + 5, right - 5
                ink = max(
                    sum(grey[yy * width + x] < 210 for x in range(inner_left, inner_right))
                    for yy in range(max(0, y - 2), min(height, y + 3))
                )
                if ink >= (inner_right - inner_left) * 0.7:
                    column_bounds.append(y)
            if column_bounds[-1] != period_bounds[-1]:
                column_bounds.append(period_bounds[-1])
            for top, bottom in zip(column_bounds, column_bounds[1:]):
                first = boundary_index[top]
                last = boundary_index[bottom] - 1
                covered = periods[first:last + 1]
                if not covered or any(not label.isdigit() or span is None for label, span in covered):
                    continue  # no lesson reminders for Recess or Lunch
                start, end = covered[0][1][0], covered[-1][1][1]
                if start >= end:
                    continue
                box = (left + 5, top + 3, right - 5, bottom - 3)
                crop_ink = image.crop(box).convert("L").tobytes()
                dark = sum(value < 150 for value in crop_ink)
                if dark < max(12, round(len(crop_ink) * 0.001)):
                    continue  # Tesseract can hallucinate letters on empty white cells
                cell_text = words_in_box(*box) or ocr_box(*box, cell_languages, 6)
                if not cell_text.strip():
                    continue
                # The Latin model is useful for class suffixes such as 1X,
                # which the Chinese model sometimes drops from a class name.
                missing_suffix = re.search(r"中[一二三四五六七八九]\s*([1-6])(?![A-Za-z0-9])", cell_text)
                if missing_suffix and header_language:
                    latin = ocr_box(*box, header_language, 6)
                    suffix = re.search(rf"{missing_suffix.group(1)}([A-Z])\b", latin, re.IGNORECASE)
                    if suffix:
                        cell_text = cell_text[:missing_suffix.end()] + suffix.group(1).upper() + cell_text[missing_suffix.end():]
                rooms = re.findall(r"(?i)\bRM\s*\d+(?:\s*,\s*RM?\s*\d+)*", cell_text)
                location = ", ".join(rooms)
                if "全校" in cell_text and not location:
                    location = "全校"
                cleaned = re.sub(r"(?i)\bRM\s*\d+(?:\s*,\s*RM?\s*\d+)*", " ", cell_text)
                if owner:
                    cleaned = cleaned.replace(owner, " ")
                cleaned = cleaned.replace("全校", " ") if location == "全校" else cleaned
                lines = []
                for line in cleaned.splitlines():
                    line = " ".join(line.split())
                    if not line or line in lines:
                        continue
                    if lines and "/" in line and not re.search(r"\d", line):
                        continue  # list of staff names, not the lesson title
                    lines.append(line)
                if (location == "全校" and len(lines) == 2
                        and all(re.fullmatch(r"[\u3400-\u9fff]{2,3}", line) for line in lines)
                        and lines[0][-1] == lines[1][-1]):
                    lines = [lines[1]]
                title = " ".join(lines)
                if not title:
                    continue
                # Crowded, overlapping print in the source is unreliable for
                # automatic reminders even when OCR returns plausible words.
                needs_review = dark / len(crop_ink) > 0.065
                if needs_review:
                    title = f"[需核對] {title}"
                rows.append([WEEKDAYS[weekday], start, end, title, location, "否" if needs_review else "是"])

    # Repeated lessons provide evidence for a few OCR typos. Reconcile only
    # within the same room and class group when a title appears twice elsewhere.
    groups: dict[tuple[str, str], list[list[object]]] = defaultdict(list)
    for row in rows[1:]:
        group = re.search(r"\b[1-6][A-Z](?:/[1-6][A-Z])+\b", str(row[3]))
        if group and row[5] == "是":
            groups[(str(row[4]), group.group())].append(row)

    def within_two_edits(left: str, right: str) -> bool:
        if left == right or abs(len(left) - len(right)) > 2:
            return False
        previous = list(range(len(right) + 1))
        for index, char in enumerate(left, start=1):
            current = [index]
            for column, other in enumerate(right, start=1):
                current.append(min(current[-1] + 1, previous[column] + 1,
                                   previous[column - 1] + (char != other)))
            if min(current) > 2:
                return False
            previous = current
        return previous[-1] <= 2

    for group_rows in groups.values():
        common, count = Counter(str(row[3]) for row in group_rows).most_common(1)[0]
        if count < 2:
            continue
        for row in group_rows:
            if within_two_edits(str(row[3]), common):
                row[3] = common
    return _rows_text(rows) if len(rows) > 1 else ""


def _ocr_grid(path: Path, tesseract: Path, languages: str) -> str:
    bounds = _grid_boundaries(path)
    if bounds is None:
        return ""
    from PIL import Image

    xs, ys = bounds
    with Image.open(path) as image, tempfile.TemporaryDirectory() as directory:
        rows = []
        temp_path = Path(directory) / "cell.png"
        for top, bottom in zip(ys, ys[1:]):
            cells = []
            for left, right in zip(xs, xs[1:]):
                cell = image.crop((left + 5, top + 5, right - 5, bottom - 5)).convert("RGB")
                if cell.convert("L").getextrema()[0] > 200:
                    cells.append("")
                    continue
                cell.save(temp_path)
                cells.append(_run_tesseract(tesseract, temp_path, languages, 6))
            rows.append(cells)
    return _rows_text(rows)


def _ocr_image(path: Path, native_text: str = "") -> str:
    tesseract = find_tesseract()
    if tesseract is None:
        raise RuntimeError("讀取圖片或掃描 PDF 需要安裝 Tesseract OCR（繁體中文及英文語言資料）")
    languages = subprocess.run(
        [str(tesseract), "--list-langs"], capture_output=True, encoding="utf-8", errors="replace",
        timeout=15, check=False,
    ).stdout
    selected = [language for language in ("chi_tra", "eng") if language in languages.splitlines()]
    if not selected:
        raise RuntimeError("Tesseract OCR 沒有繁體中文或英文語言資料")
    language_arg = "+".join(selected)
    weekly_text = _weekly_grid_text(path, native_text=native_text)
    if weekly_text:
        return weekly_text
    grid_text = _ocr_grid(path, tesseract, language_arg)
    if grid_text and parse_timetable_text(grid_text):
        return grid_text
    text = _run_tesseract(tesseract, path, language_arg, 6)
    if not text:
        text = _run_tesseract(tesseract, path, language_arg, 11)
    return text or grid_text


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    rendered = None
    try:
        if reader.is_encrypted:
            raise ValueError("加密 PDF 無法讀取")
        if len(reader.pages) > 12:
            raise ValueError("時間表 PDF 超過 12 頁，請先用 PDF 工具擷取相關頁面")
        texts = []
        native_texts = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text(extraction_mode="layout") or ""
            native_text = page.extract_text() or ""
            native_texts.append(native_text)
            if len(text.strip()) < 20:
                if rendered is None:
                    import pypdfium2 as pdfium

                    rendered = pdfium.PdfDocument(str(path))
                with tempfile.TemporaryDirectory() as directory:
                    image_path = Path(directory) / "page.png"
                    rendered[index].render(scale=3).to_pil().save(image_path)
                    text = _ocr_image(image_path, native_text=native_text)
            texts.append(text)
        extracted = "\n\n".join(texts)
        if parse_timetable_text(extracted) or find_tesseract() is None:
            return extracted
        if rendered is None:
            import pypdfium2 as pdfium

            rendered = pdfium.PdfDocument(str(path))
        ocr_pages = []
        for index in range(len(reader.pages)):
            with tempfile.TemporaryDirectory() as directory:
                image_path = Path(directory) / "page.png"
                rendered[index].render(scale=3).to_pil().save(image_path)
                ocr_pages.append(_ocr_image(image_path, native_text=native_texts[index]))
        recognized = "\n\n".join(ocr_pages)
        return recognized if parse_timetable_text(recognized) else extracted + "\n\n" + recognized
    finally:
        if rendered is not None:
            rendered.close()
        reader.close()


def _vision_rows(path: Path, model: str, method: str, api_key: str) -> str:
    def recognize(image_path: Path) -> list[dict[str, object]]:
        if method == "openrouter":
            return recognize_image_openrouter(image_path, api_key, model)
        return recognize_image(image_path, model)

    if path.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium

        reader = PdfReader(str(path))
        try:
            if reader.is_encrypted:
                raise ValueError("加密 PDF 無法讀取")
            if len(reader.pages) > 12:
                raise ValueError("時間表 PDF 超過 12 頁，請先用 PDF 工具擷取相關頁面")
        finally:
            reader.close()
        document = pdfium.PdfDocument(str(path))
        try:
            with tempfile.TemporaryDirectory() as directory:
                entries = []
                for index in range(len(document)):
                    image_path = Path(directory) / f"page-{index + 1}.png"
                    document[index].render(scale=2.5).to_pil().save(image_path)
                    entries.extend(recognize(image_path))
        finally:
            document.close()
    else:
        entries = recognize(path)

    rows: list[list[object]] = [["星期", "開始", "結束", "課堂", "地點"]]
    for entry in entries:
        weekday = entry.get("weekday")
        start = entry.get("start_time")
        end = entry.get("end_time")
        title = entry.get("title")
        location = entry.get("location", "")
        if not (type(weekday) is int and 0 <= weekday <= 6 and
                isinstance(start, str) and isinstance(end, str) and
                isinstance(title, str) and isinstance(location, str)):
            continue
        start_time, end_time = _clock(start), _clock(end)
        if start_time is None or end_time is None or start_time >= end_time or not title.strip():
            continue
        rows.append([WEEKDAYS[weekday], start_time, end_time, title.strip(), location.strip()])
    return _rows_text(rows)


def _nvidia_image_text(path: Path, api_key: str, native_text: str = "") -> str:
    """Place OCR words back in their timetable cells before parsing."""
    from PIL import Image, ImageOps

    with tempfile.TemporaryDirectory() as directory:
        normalized = Path(directory) / "page.png"
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
            image.thumbnail((3200, 3200), Image.Resampling.LANCZOS)
            width, height = image.size
            image.save(normalized)
        words = recognize_image_nvidia(normalized, api_key)
        if not words:
            raise RuntimeError("NVIDIA OCR 沒有辨識到文字，請檢查圖片清晰度")
        weekly_text = _weekly_grid_text(normalized, words=words, native_text=native_text)
        if weekly_text:
            return weekly_text
        bounds = _grid_boundaries(normalized)

    if bounds:
        xs, ys = bounds
        cells: list[list[list[dict[str, object]]]] = [
            [[] for _ in range(len(xs) - 1)] for _ in range(len(ys) - 1)
        ]
        for word in words:
            column = bisect_right(xs, float(word["x"]) * width) - 1
            row = bisect_right(ys, float(word["y"]) * height) - 1
            if 0 <= row < len(cells) and 0 <= column < len(cells[row]):
                cells[row][column].append(word)
        rows = [
            [" ".join(str(word["text"]) for word in sorted(cell, key=lambda item: (round(float(item["y"]) / 0.012), float(item["x"]))))
             for cell in row]
            for row in cells
        ]
        grid_text = _rows_text(rows)
        if parse_timetable_text(grid_text):
            return grid_text

    positioned_text = _positioned_ocr_grid(words)
    if positioned_text:
        return positioned_text

    lines: list[list[dict[str, object]]] = []
    for word in sorted(words, key=lambda item: (float(item["y"]), float(item["x"]))):
        if not lines or abs(float(word["y"]) - float(lines[-1][0]["y"])) > 0.012:
            lines.append([])
        lines[-1].append(word)
    return "\n".join(
        " ".join(str(word["text"]) for word in sorted(line, key=lambda item: float(item["x"])))
        for line in lines
    )


def _positioned_ocr_grid(words: list[dict[str, object]]) -> str:
    """Rebuild an unruled weekly grid from weekday headings and printed times."""
    headings = sorted(
        ((float(word["x"]), float(word["y"]), day)
         for word in words if (day := _weekday(str(word["text"]))) is not None),
        key=lambda item: (item[1], item[0]),
    )
    groups: list[list[tuple[float, float, int]]] = []
    for heading in headings:
        if not groups or heading[1] - groups[-1][0][1] > 0.025:
            groups.append([])
        groups[-1].append(heading)
    candidates = [group for group in groups if len({item[2] for item in group}) >= 2]
    if not candidates:
        return ""
    header = sorted(max(candidates, key=lambda group: (len({item[2] for item in group}), -group[0][1])),
                    key=lambda item: item[0])
    day_centers = [item[0] for item in header]
    if len(day_centers) < 2 or day_centers[-1] - day_centers[0] < 0.12:
        return ""
    header_y = sum(item[1] for item in header) / len(header)
    first_boundary = max(0.0, day_centers[0] - (day_centers[1] - day_centers[0]) / 2)

    def lines(items: list[dict[str, object]]) -> list[list[dict[str, object]]]:
        grouped: list[list[dict[str, object]]] = []
        for item in sorted(items, key=lambda word: (float(word["y"]), float(word["x"]))):
            if not grouped or float(item["y"]) - float(grouped[-1][0]["y"]) > 0.015:
                grouped.append([])
            grouped[-1].append(item)
        return grouped

    time_words = [word for word in words if float(word["x"]) < first_boundary
                  and float(word["y"]) > header_y + 0.02]
    row_bands: list[tuple[float, float, str]] = []
    # aSc-style timetables print the period number and class above the time.
    # Period numbers give more accurate row boundaries than the time's position.
    markers = sorted(
        (float(word["y"]) for word in time_words
         if re.fullmatch(r"\d{1,2}", str(word["text"]).strip())
         and 1 <= int(str(word["text"]).strip()) <= 20),
    )
    markers = [y for index, y in enumerate(markers) if index == 0 or y - markers[index - 1] > 0.02]
    if len(markers) >= 2:
        for index, marker_y in enumerate(markers):
            before = max(header_y + 0.015, marker_y - 0.015)
            after = markers[index + 1] - 0.015 if index + 1 < len(markers) else min(
                1.0, marker_y + (marker_y - markers[index - 1]) - 0.015
            )
            period_words = [word for word in time_words if before <= float(word["y"]) < after]
            period = " ".join(
                str(word["text"]) for line in lines(period_words)
                for word in sorted(line, key=lambda item: float(item["x"]))
            )
            if _span(period):
                row_bands.append((before, after, period))

    if not row_bands:
        time_rows: list[tuple[float, str]] = []
        for line in lines(time_words):
            text = " ".join(str(word["text"]) for word in sorted(line, key=lambda word: float(word["x"])))
            if _span(text):
                time_rows.append((sum(float(word["y"]) for word in line) / len(line), text))
        for index, (center_y, period) in enumerate(time_rows):
            before = (time_rows[index - 1][0] + center_y) / 2 if index else header_y + 0.015
            after = (center_y + time_rows[index + 1][0]) / 2 if index + 1 < len(time_rows) else (
                center_y + (center_y - time_rows[index - 1][0]) / 2 if index else center_y + 0.08
            )
            row_bands.append((before, after, period))
    if not row_bands:
        return ""

    column_boundaries = [(left + right) / 2 for left, right in zip(day_centers, day_centers[1:])]
    rows: list[list[object]] = [["時間", *[WEEKDAYS[day] for _, _, day in header]]]
    for before, after, period in row_bands:
        cells: list[list[dict[str, object]]] = [[] for _ in header]
        for word in words:
            x, y = float(word["x"]), float(word["y"])
            if not before <= y < after or x < first_boundary:
                continue
            column = bisect_right(column_boundaries, x)
            cells[column].append(word)
        row: list[object] = [period]
        for cell in cells:
            row.append("\n".join(
                " ".join(str(word["text"]) for word in sorted(line, key=lambda word: float(word["x"])))
                for line in lines(cell)
            ))
        rows.append(row)
    text = _rows_text(rows)
    return text if parse_timetable_text(text) else ""


def _nvidia_text(path: Path, api_key: str) -> str:
    if path.suffix.lower() != ".pdf":
        return _nvidia_image_text(path, api_key)
    import pypdfium2 as pdfium

    reader = PdfReader(str(path))
    try:
        if reader.is_encrypted:
            raise ValueError("加密 PDF 無法讀取")
        if len(reader.pages) > 12:
            raise ValueError("時間表 PDF 超過 12 頁，請先用 PDF 工具擷取相關頁面")
        native_pages = [page.extract_text() or "" for page in reader.pages]
    finally:
        reader.close()
    document = pdfium.PdfDocument(str(path))
    try:
        with tempfile.TemporaryDirectory() as directory:
            pages = []
            for index in range(len(document)):
                image_path = Path(directory) / f"page-{index + 1}.png"
                document[index].render(scale=3).to_pil().save(image_path)
                pages.append(_nvidia_image_text(image_path, api_key, native_pages[index]))
            return "\n\n".join(pages)
    finally:
        document.close()


def extract_timetable_text(
    path: Path, method: str = "standard", model: str | None = None, api_key: str = "",
) -> str:
    suffix = path.suffix.lower()
    if method not in {"standard", "vision", "openrouter", "nvidia"}:
        raise ValueError(f"不支援的辨識方式：{method}")
    if method == "nvidia" and (suffix in IMAGE_EXTENSIONS or suffix == ".pdf"):
        return _nvidia_text(path, api_key)
    if method in {"vision", "openrouter"} and (suffix in IMAGE_EXTENSIONS or suffix == ".pdf"):
        selected_model = model if model is not None else (
            DEFAULT_OPENROUTER_MODEL if method == "openrouter" else DEFAULT_MODEL
        )
        return _vision_rows(path, selected_model, method, api_key)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(path)
    if suffix == ".xls":
        return _read_xls(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".doc":
        return _read_doc(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in IMAGE_EXTENSIONS:
        return _ocr_image(path)
    raise ValueError(f"不支援的時間表格式：{suffix}")


def _weekday(value: str) -> int | None:
    cleaned = value.strip().lower().replace(" ", "")
    if cleaned in DAY_ALIASES:
        return DAY_ALIASES[cleaned]
    if cleaned.startswith(("星期", "週", "周")):
        return DAY_ALIASES.get(cleaned[-1])
    return None


def _clock(value: str) -> str | None:
    match = TIME_PATTERN.fullmatch(value.strip())
    if not match:
        return None
    hour, minute = int(match.group(1)), int(match.group(2))
    return f"{hour:02d}:{minute:02d}" if hour < 24 else None


def _span(value: str) -> tuple[str, str] | None:
    cleaned = re.sub(r"[:：]{2,}", ":", value)
    cleaned = re.sub(r"([-–—~～])\s*:(?=\s*[0-2]?\d\s*[:：.])", r"\1", cleaned)
    cleaned = re.sub(r"([-–—~～]){2,}", r"\1", cleaned)
    match = SPAN_PATTERN.search(cleaned)
    if not match:
        return None
    start, end = _clock(match.group(1)), _clock(match.group(2))
    if start and end and start < end:
        return start, end
    return None


def _header_kind(value: str) -> str:
    text = value.strip().lower().replace(" ", "")
    if text in {"星期", "週", "周", "日期", "day", "weekday"}:
        return "day"
    if text in {"開始", "開始時間", "上課", "start", "starttime"}:
        return "start"
    if text in {"結束", "結束時間", "下課", "end", "endtime"}:
        return "end"
    if text in {"時間", "時段", "節次時間", "time", "period"}:
        return "period"
    if text in {"科目", "課堂", "課程", "內容", "subject", "lesson", "title"}:
        return "title"
    if text in {"班別", "班級", "class", "group"}:
        return "class"
    if text in {"課室", "教室", "地點", "room", "location"}:
        return "location"
    if text in {"啟用", "enabled"}:
        return "enabled"
    return ""


def _table_entries(rows: list[list[str]]) -> list[SuggestedEntry]:
    for header_index, header in enumerate(rows[:5]):
        fields = {_header_kind(cell): index for index, cell in enumerate(header) if _header_kind(cell)}
        if "day" in fields and ("period" in fields or {"start", "end"} <= fields.keys()):
            entries = []
            for row in rows[header_index + 1:]:
                def value(key: str) -> str:
                    index = fields.get(key, -1)
                    return row[index].strip() if 0 <= index < len(row) else ""

                weekday = _weekday(value("day"))
                times = _span(value("period")) if "period" in fields else None
                if times is None and "start" in fields and "end" in fields:
                    start, end = _clock(value("start")), _clock(value("end"))
                    times = (start, end) if start and end and start < end else None
                title = " ".join(part for part in (value("title"), value("class")) if part)
                if weekday is not None and times and title:
                    enabled = value("enabled").lower() not in {"0", "否", "false", "no"}
                    entries.append(SuggestedEntry(weekday, *times, title, value("location"), enabled))
            return entries

        day_columns = {index: day for index, cell in enumerate(header) if (day := _weekday(cell)) is not None}
        if len(day_columns) >= 2:
            # OCR sometimes reads adjacent weekday headers as the same character.
            previous_day: int | None = None
            for index in sorted(day_columns):
                if previous_day is not None and day_columns[index] == previous_day and previous_day < 6:
                    day_columns[index] = previous_day + 1
                previous_day = day_columns[index]
            entries = []
            for row in rows[header_index + 1:]:
                period = next((_span(cell) for index, cell in enumerate(row) if index not in day_columns and _span(cell)), None)
                if not period:
                    start_column = next((index for index, cell in enumerate(header) if _header_kind(cell) == "start"), None)
                    end_column = next((index for index, cell in enumerate(header) if _header_kind(cell) == "end"), None)
                    if start_column is not None and end_column is not None:
                        start = _clock(row[start_column]) if start_column < len(row) else None
                        end = _clock(row[end_column]) if end_column < len(row) else None
                        period = (start, end) if start and end and start < end else None
                if not period:
                    times = [
                        _clock(match.group())
                        for index, cell in enumerate(row) if index not in day_columns
                        for match in TIME_PATTERN.finditer(cell)
                    ]
                    if len(times) == 2 and times[0] and times[1] and times[0] < times[1]:
                        period = (times[0], times[1])
                if not period:
                    continue
                for index, weekday in day_columns.items():
                    if index < len(row) and row[index].strip():
                        parts = [part.strip() for part in row[index].splitlines() if part.strip()]
                        title = " ".join(parts[:1])
                        location = " ".join(parts[1:])
                        entries.append(SuggestedEntry(weekday, *period, title, location))
            return entries
    return []


def _line_entry(line: str) -> SuggestedEntry | None:
    day_match = DAY_PATTERN.search(line)
    period = _span(line)
    if not day_match or not period:
        return None
    weekday = _weekday(day_match.group())
    if weekday is None:
        return None
    title = DAY_PATTERN.sub("", line, count=1)
    title = SPAN_PATTERN.sub("", title, count=1).strip(" \t:：-–—|,，")
    title = " ".join(title.split())
    return SuggestedEntry(weekday, *period, title) if title else None


def parse_timetable_text(text: str) -> list[SuggestedEntry]:
    entries: list[SuggestedEntry] = []
    for block in re.split(r"\n\s*\n", text.strip()):
        if not block.strip():
            continue
        rows = list(csv.reader(io.StringIO(block), delimiter="\t"))
        parsed = _table_entries(rows)
        if parsed:
            entries.extend(parsed)
        else:
            entries.extend(entry for line in block.splitlines() if (entry := _line_entry(line)))
    return sorted(set(entries), key=lambda item: (item.weekday, item.start_time, item.title))
