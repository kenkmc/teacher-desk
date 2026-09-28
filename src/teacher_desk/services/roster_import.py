from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


HEADER_MAP = {
    "學號": "student_no",
    "編號": "student_no",
    "座號": "student_no",
    "student_no": "student_no",
    "id": "student_no",
    "中文姓名": "name",
    "中文名": "name",
    "姓名": "name",
    "名字": "name",
    "chinese_name": "name",
    "chinese": "name",
    "name": "name",
    "英文姓名": "english_name",
    "英文名": "english_name",
    "english_name": "english_name",
    "english": "english_name",
    "班別": "class_name",
    "班級": "class_name",
    "班": "class_name",
    "class": "class_name",
    "class_name": "class_name",
    "class_code": "class_name",
}


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower()


def parse_table_rows(headers: list[str], rows: list[list[object]]) -> list[tuple[str, str, str, str]]:
    mapped = [_normalize_header(h) for h in headers]
    keys: dict[str, int] = {}
    for index, header in enumerate(mapped):
        field = HEADER_MAP.get(header) or HEADER_MAP.get(headers[index].strip() if index < len(headers) else "")
        if field and field not in keys:
            keys[field] = index
    if "name" not in keys and "english_name" not in keys:
        raise ValueError("找不到「中文姓名」或「英文姓名」欄")

    students: list[tuple[str, str, str, str]] = []
    for row_index, row in enumerate(rows, start=1):
        name = ""
        if "name" in keys and keys["name"] < len(row):
            name = str(row[keys["name"]] or "").strip()
        english = ""
        if "english_name" in keys and keys["english_name"] < len(row):
            english = str(row[keys["english_name"]] or "").strip()
        if not name and not english:
            continue
        if "student_no" in keys:
            student_no = str(row[keys["student_no"]] if keys["student_no"] < len(row) else "").strip()
        else:
            student_no = str(row_index)
        class_name = ""
        if "class_name" in keys and keys["class_name"] < len(row):
            class_name = str(row[keys["class_name"]] or "").strip()
        students.append((student_no or str(row_index), name, english, class_name))
    if not students:
        raise ValueError("沒有可匯入的學生列")
    return students


def import_students(path: Path) -> list[tuple[str, str, str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            raise ValueError("試算表是空的")
        headers = [str(cell or "") for cell in rows[0]]
        body = [list(row) for row in rows[1:]]
        return parse_table_rows(headers, body)
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            first_line = handle.readline()
            handle.seek(0)
            delimiter = "\t" if first_line.count("\t") >= first_line.count(",") and "\t" in first_line else ","
            parsed = [row for row in csv.reader(handle, delimiter=delimiter, strict=True) if any(cell.strip() for cell in row)]
        if not parsed:
            raise ValueError("CSV 是空的")
        return parse_table_rows(parsed[0], parsed[1:])
    raise ValueError(f"不支援的檔案格式：{suffix}")
