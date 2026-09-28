import io
import json
import sys
from datetime import datetime, time
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication
from reportlab.pdfgen import canvas
import pytest

from teacher_desk.db import Database, TimetableEntry
from teacher_desk.services.reminders import due_occurrence
from teacher_desk.services.timetable_import import extract_timetable_text, find_tesseract, parse_timetable_text
from teacher_desk.services.timetable_vision import recognize_image, recognize_image_nvidia, recognize_image_openrouter


def test_bundled_tesseract_is_preferred(tmp_path: Path, monkeypatch) -> None:
    bundled = tmp_path / "tesseract" / "tesseract.exe"
    bundled.parent.mkdir()
    bundled.touch()
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    assert find_tesseract() == bundled


def test_excel_list_import_and_persist_reminder(tmp_path: Path) -> None:
    path = tmp_path / "weekly.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["星期", "開始", "結束", "科目", "班別", "課室"])
    sheet.append(["星期一", time(8, 30), time(9, 15), "Math", "1A", "201"])
    workbook.save(path)

    suggested = parse_timetable_text(extract_timetable_text(path))
    assert [(entry.weekday, entry.start_time, entry.end_time, entry.title, entry.location) for entry in suggested] == [
        (0, "08:30", "09:15", "Math 1A", "201")
    ]

    db = Database(tmp_path / "data.sqlite3")
    db.replace_timetable_entries([
        TimetableEntry(None, item.weekday, item.start_time, item.end_time, item.title, item.location, 10)
        for item in suggested
    ])
    saved = db.list_timetable_entries()[0]
    assert due_occurrence(saved, datetime(2026, 9, 28, 8, 20, 15)) == "2026-09-28"
    assert db.mark_reminder_delivered(saved.id, "2026-09-28") is True
    assert db.mark_reminder_delivered(saved.id, "2026-09-28") is False
    db.replace_timetable_entries([TimetableEntry(None, 0, "08:30", "09:15", "Math 1A", "202", 10)])
    assert db.list_timetable_entries()[0].id == saved.id
    assert db.mark_reminder_delivered(saved.id, "2026-09-28") is False


def test_excel_grid_import(tmp_path: Path) -> None:
    path = tmp_path / "grid.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["時間", "星期一", "星期二"])
    sheet.append(["08:30-09:15", "Math\n201室", "English"])
    workbook.save(path)

    suggested = parse_timetable_text(extract_timetable_text(path))
    assert [(entry.weekday, entry.title, entry.location) for entry in suggested] == [
        (0, "Math", "201室"),
        (1, "English", ""),
    ]


def test_word_table_and_pdf_text_import(tmp_path: Path) -> None:
    word_path = tmp_path / "weekly.docx"
    document = Document()
    table = document.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ["星期", "時間", "課堂", "地點"]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["星期三", "10:00-10:45", "Science", "Lab"]):
        cell.text = value
    document.save(word_path)
    assert parse_timetable_text(extract_timetable_text(word_path))[0].title == "Science"

    pdf_path = tmp_path / "weekly.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(40, 750, "Thursday 13:30-14:15 Geography Room 2")
    pdf.save()
    parsed = parse_timetable_text(extract_timetable_text(pdf_path))
    assert [(entry.weekday, entry.start_time, entry.title) for entry in parsed] == [
        (3, "13:30", "Geography Room 2")
    ]


def test_reminder_can_start_on_previous_day() -> None:
    entry = TimetableEntry(1, 0, "00:30", "01:15", "Early class", remind_before=60)
    assert due_occurrence(entry, datetime(2026, 9, 27, 23, 30, 20)) == "2026-09-28"
    assert due_occurrence(entry, datetime(2026, 9, 27, 23, 35)) is None


def test_grid_ocr_duplicate_weekday_header_is_reconciled() -> None:
    text = "時間\t星期一\t星期二\t星期二\n08:30-09:15\t中文\t數學\t英文"
    assert [entry.weekday for entry in parse_timetable_text(text)] == [0, 1, 2]


def test_grid_with_separate_start_and_end_columns() -> None:
    text = "開始\t結束\t星期一\t星期二\n08:30\t09:15\t數學\t英文"
    assert [(entry.weekday, entry.start_time, entry.end_time, entry.title)
            for entry in parse_timetable_text(text)] == [
        (0, "08:30", "09:15", "數學"), (1, "08:30", "09:15", "英文"),
    ]


def test_import_auto_populates_timetable_and_reminder_draft(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    from teacher_desk.modules.timetable.widget import TimetableWidget

    database = Database(tmp_path / "draft.sqlite3")
    widget = TimetableWidget(database)
    widget._imported("時間\t星期一\t星期二\n08:30-09:15\tMath\tEnglish")
    assert widget.table.rowCount() == 2
    assert [widget.table.item(row, 3).text() for row in range(2)] == ["Math", "English"]
    assert [widget.table.item(row, 5).text() for row in range(2)] == ["10", "10"]
    assert all(widget.table.item(row, 6).checkState() == Qt.CheckState.Checked for row in range(2))
    assert database.list_timetable_entries() == []
    widget._imported("星期\t開始\t結束\t課堂\t地點\t啟用\n星期三\t15:30\t16:05\t[需核對] 文字擠疊\t\t否")
    assert widget.table.rowCount() == 1
    assert widget.table.item(0, 6).checkState() == Qt.CheckState.Unchecked
    widget.close()
    assert app is not None


@pytest.mark.skipif(find_tesseract() is None, reason="Tesseract not installed")
def test_ruled_pdf_and_image_merge_lesson_periods(tmp_path: Path, monkeypatch) -> None:
    import pypdfium2 as pdfium

    pdf_path = tmp_path / "weekly-grid.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(600, 420))
    pdf.setLineWidth(1)
    xs = (20, 110, 260, 410, 560)
    for x in xs:
        pdf.line(x, 420 - 20, x, 420 - 380)
    for y in (20, 80, 200, 260, 320, 380):
        pdf.line(xs[0], 420 - y, xs[-1], 420 - y)
    for left, right in ((20, 110), (260, 410), (410, 560)):
        pdf.line(left, 420 - 140, right, 420 - 140)
    pdf.setFont("Helvetica", 15)
    for label, x in (("Mo", 165), ("Tu", 315), ("We", 465)):
        pdf.drawString(x, 420 - 60, label)
    pdf.setFont("Helvetica", 11)
    for label, start, end, top in (
        ("1", "08:00", "08:40", 95), ("2", "08:40", "09:20", 155),
        ("Recess1", "09:20", "09:40", 215), ("3", "09:40", "10:20", 275),
        ("4", "10:20", "11:00", 335),
    ):
        pdf.drawString(32, 420 - top, label)
        pdf.drawString(27, 420 - (top + 25), f"{start} - {end}")
    pdf.drawString(125, 420 - 105, "Math RM101")
    pdf.drawString(275, 420 - 105, "English RM102")
    pdf.drawString(250, 420 - 230, "Recess1")
    pdf.drawString(425, 420 - 285, "Science RM103")
    pdf.save()

    document = pdfium.PdfDocument(str(pdf_path))
    image_path = tmp_path / "weekly-grid.png"
    document[0].render(scale=3).to_pil().save(image_path)
    document.close()

    for path in (pdf_path, image_path):
        entries = parse_timetable_text(extract_timetable_text(path))
        assert [(entry.weekday, entry.start_time, entry.end_time, entry.title, entry.location)
                for entry in entries] == [
            (0, "08:00", "09:20", "Math", "RM101"),
            (1, "08:00", "08:40", "English", "RM102"),
            (2, "09:40", "10:20", "Science", "RM103"),
        ]

    words = [
        {"text": text, "x": x / 1800, "y": y / 1260}
        for text, x, y in (
            ("Mo", 495, 180), ("Tu", 945, 180), ("We", 1395, 180),
            ("Math", 400, 320), ("RM101", 570, 320),
            ("English", 850, 320), ("RM102", 1020, 320),
            ("Science", 1300, 850), ("RM103", 1470, 850),
        )
    ]
    monkeypatch.setattr("teacher_desk.services.timetable_import.recognize_image_nvidia", lambda *_: words)
    cloud_entries = parse_timetable_text(extract_timetable_text(image_path, "nvidia", api_key="test-key"))
    assert [(entry.weekday, entry.start_time, entry.end_time, entry.title) for entry in cloud_entries] == [
        (0, "08:00", "09:20", "Math"),
        (1, "08:00", "08:40", "English"),
        (2, "09:40", "10:20", "Science"),
    ]


def test_local_vision_result_is_validated_before_preview(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "timetable.png"
    Image.new("RGB", (100, 100), "white").save(image_path)
    monkeypatch.setattr("teacher_desk.services.timetable_import.recognize_image", lambda _path, _model: [
        {"weekday": 0, "start_time": "08:30", "end_time": "09:15", "title": "數學 1A", "location": "201室"},
        {"weekday": 8, "start_time": "10:00", "end_time": "10:45", "title": "錯誤", "location": ""},
        {"weekday": 1, "start_time": "不明", "end_time": "11:00", "title": "猜測", "location": ""},
    ])

    suggested = parse_timetable_text(extract_timetable_text(image_path, "vision"))
    assert [(item.weekday, item.start_time, item.title, item.location) for item in suggested] == [
        (0, "08:30", "數學 1A", "201室")
    ]


def test_local_vision_request_uses_loopback_and_schema(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "weekly.png"
    Image.new("RGB", (50, 50), "white").save(image_path)
    sent = {}

    class FakeOpener:
        def open(self, request, timeout):
            sent["url"] = request.full_url
            sent["body"] = json.loads(request.data)
            sent["timeout"] = timeout
            reply = {"message": {"content": json.dumps({"entries": [{
                "weekday": 2, "start_time": "09:00", "end_time": "09:45",
                "title": "英文 2B", "location": "302室",
            }]})}}
            return io.BytesIO(json.dumps(reply).encode("utf-8"))

    monkeypatch.setattr("teacher_desk.services.timetable_vision.urllib.request.build_opener", lambda *_: FakeOpener())
    assert recognize_image(image_path)[0]["title"] == "英文 2B"
    assert sent["url"] == "http://127.0.0.1:11434/api/chat"
    assert sent["body"]["model"] == "qwen3-vl:2b"
    assert sent["body"]["format"]["required"] == ["entries"]
    assert sent["body"]["messages"][0]["images"]
    assert sent["body"]["stream"] is False


def test_openrouter_free_vision_request_and_preview(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "weekly.png"
    Image.new("RGB", (50, 50), "white").save(image_path)
    sent = {}

    def fake_urlopen(request, timeout):
        sent["url"] = request.full_url
        sent["key"] = request.get_header("Authorization")
        sent["body"] = json.loads(request.data)
        reply = {"choices": [{"message": {"content": json.dumps({"entries": [{
            "weekday": 4, "start_time": "13:00", "end_time": "13:45",
            "title": "中文 2A", "location": "101室",
        }]})}}]}
        return io.BytesIO(json.dumps(reply).encode("utf-8"))

    monkeypatch.setattr("teacher_desk.services.timetable_vision.urllib.request.urlopen", fake_urlopen)
    suggested = parse_timetable_text(extract_timetable_text(image_path, "openrouter", api_key="test-key"))
    assert [(item.weekday, item.start_time, item.title) for item in suggested] == [(4, "13:00", "中文 2A")]
    assert sent["url"] == "https://openrouter.ai/api/v1/chat/completions"
    assert sent["key"] == "Bearer test-key"
    assert sent["body"]["model"] == "google/gemma-4-31b-it:free"
    assert sent["body"]["response_format"] == {"type": "json_object"}
    image_data = sent["body"]["messages"][0]["content"][1]["image_url"]["url"]
    assert image_data.startswith("data:image/png;base64,")

    with pytest.raises(ValueError, match="免費模型"):
        recognize_image_openrouter(image_path, "test-key", "google/gemma-4-31b-it")


def test_nvidia_ocr_request_and_grid_preview(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "grid.png"
    image = Image.new("RGB", (300, 200), "white")
    from PIL import ImageDraw
    draw = ImageDraw.Draw(image)
    for x in (0, 100, 200, 299):
        draw.line((x, 0, x, 199), fill="black", width=2)
    for y in (0, 100, 199):
        draw.line((0, y, 299, y), fill="black", width=2)
    image.save(image_path)

    samples = [
        ("時間", .16, .25), ("星期一", .50, .25), ("星期二", .83, .25),
        ("08:30-09:15", .16, .75), ("Math", .50, .75), ("Science", .83, .75),
    ]
    detections = [{
        "text_prediction": {"text": word, "confidence": .99},
        "bounding_box": {"points": [{"x": x, "y": y} for _ in range(4)]},
    } for word, x, y in samples]
    sent = {}

    class FakeResponse(io.BytesIO):
        status = 200
        headers = {}

    def fake_urlopen(request, timeout):
        sent["url"] = request.full_url
        sent["key"] = request.get_header("Authorization")
        sent["body"] = json.loads(request.data)
        return FakeResponse(json.dumps({"data": [{"text_detections": detections}]}).encode())

    monkeypatch.setattr("teacher_desk.services.timetable_vision.urllib.request.urlopen", fake_urlopen)
    text = extract_timetable_text(image_path, "nvidia", api_key="test-key")
    suggested = parse_timetable_text(text)
    assert [(item.weekday, item.start_time, item.title) for item in suggested] == [
        (0, "08:30", "Math"), (1, "08:30", "Science"),
    ]
    assert sent["url"] == "https://ai.api.nvidia.com/v1/cv/nvidia/nemotron-ocr-v2"
    assert sent["key"] == "Bearer test-key"
    assert sent["body"]["merge_levels"] == ["word"]
    assert sent["body"]["input"][0]["url"].startswith("data:image/png;base64,")


def test_nvidia_ocr_polls_async_result(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "weekly.png"
    Image.new("RGB", (50, 50), "white").save(image_path)
    urls = []

    class FakeResponse(io.BytesIO):
        def __init__(self, status, headers, body=b""):
            super().__init__(body)
            self.status = status
            self.headers = headers

    def fake_urlopen(request, timeout):
        urls.append(request.full_url)
        if len(urls) == 1:
            return FakeResponse(202, {"NVCF-REQID": "request-123"})
        result = {"data": [{"text_detections": [{
            "text_prediction": {"text": "Math", "confidence": .9},
            "bounding_box": {"points": [{"x": .5, "y": .5}]},
        }]}]}
        return FakeResponse(200, {}, json.dumps(result).encode())

    monkeypatch.setattr("teacher_desk.services.timetable_vision.urllib.request.urlopen", fake_urlopen)
    monkeypatch.setattr("teacher_desk.services.timetable_vision.time.sleep", lambda _: None)
    assert recognize_image_nvidia(image_path, "test-key")[0]["text"] == "Math"
    assert urls[1] == "https://api.nvcf.nvidia.com/v2/nvcf/pexec/status/request-123"


def test_nvidia_unruled_grid_creates_entries(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "unruled.png"
    Image.new("RGB", (1000, 700), "white").save(image_path)
    words = [
        {"text": "時間", "x": .12, "y": .18},
        {"text": "星期一", "x": .40, "y": .18},
        {"text": "星期二", "x": .72, "y": .18},
        {"text": "08:30-09:15", "x": .12, "y": .36},
        {"text": "Math", "x": .40, "y": .36},
        {"text": "English", "x": .72, "y": .36},
        {"text": "09:20-10:05", "x": .12, "y": .55},
        {"text": "Science", "x": .40, "y": .55},
    ]
    monkeypatch.setattr("teacher_desk.services.timetable_import.recognize_image_nvidia", lambda *_: words)
    suggested = parse_timetable_text(extract_timetable_text(image_path, "nvidia", api_key="test-key"))
    assert [(entry.weekday, entry.start_time, entry.title) for entry in suggested] == [
        (0, "08:30", "Math"), (0, "09:20", "Science"), (1, "08:30", "English"),
    ]


def test_nvidia_asc_layout_with_short_weekdays_and_time_below_class(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "school-schedule.png"
    Image.new("RGB", (1000, 700), "white").save(image_path)
    words = [
        {"text": day, "x": x, "y": .15}
        for day, x in zip(("Mo", "Tu", "We", "Th", "Fr"), (.30, .43, .56, .69, .82))
    ] + [
        {"text": "1", "x": .07, "y": .25},
        {"text": "中文 1X", "x": .30, "y": .25},
        {"text": "電腦 RM403", "x": .56, "y": .25},
        {"text": "8:20-8:55", "x": .09, "y": .31},
        {"text": "古敏聰", "x": .30, "y": .32},
        {"text": "2", "x": .07, "y": .40},
        {"text": "數學 2A", "x": .43, "y": .40},
        {"text": "8:55 -:9:00", "x": .09, "y": .47},
        {"text": "3", "x": .07, "y": .55},
        {"text": "Science", "x": .69, "y": .55},
        {"text": "9:30-9:50", "x": .09, "y": .61},
    ]
    monkeypatch.setattr("teacher_desk.services.timetable_import.recognize_image_nvidia", lambda *_: words)
    suggested = parse_timetable_text(extract_timetable_text(image_path, "nvidia", api_key="test-key"))
    assert [(entry.weekday, entry.start_time, entry.end_time) for entry in suggested] == [
        (0, "08:20", "08:55"), (1, "08:55", "09:00"),
        (2, "08:20", "08:55"), (3, "09:30", "09:50"),
    ]
